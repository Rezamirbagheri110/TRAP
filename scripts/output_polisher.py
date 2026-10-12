from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from statistics import median
from typing import Any, Dict, List

import sys
import os
sys.path.insert(0, str(Path(__file__).parent.parent))
from utils.schemas import normalise_iteration, normalise_behavior


def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None:
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _get_iterations(results: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Return normalised iteration list regardless of raw field name."""
    raw = results.get("iterations") or results.get("attempts") or []
    if not isinstance(raw, list):
        return []
    return [normalise_iteration(it) for it in raw if isinstance(it, dict)]


def _write_text_summary(output_dir: Path, results: Dict[str, Any], iterations: List[Dict[str, Any]]) -> Path:
    summary_path = output_dir / "summary.md"
    success = bool(results.get("success", False))
    best_score = _to_float(results.get("best_score", 0.0))
    duration = _to_float(results.get("duration_seconds", 0.0))

    scores = [_to_float(it.get("score", 0.0)) for it in iterations]
    avg_score = sum(scores) / len(scores) if scores else 0.0
    max_score = max(scores) if scores else best_score
    min_score = min(scores) if scores else 0.0
    med_score = median(scores) if scores else 0.0
    iter_secs = [_to_float(it.get("iteration_seconds", 0.0)) for it in iterations]
    avg_iter_sec = sum(iter_secs) / len(iter_secs) if iter_secs else 0.0
    source_counts = Counter(str(it.get("source", "unknown")) for it in iterations)
    top_sources = source_counts.most_common(5)

    lines = [
        "# Improved Output Summary",
        "",
        f"- Generated: `{datetime.now().isoformat(timespec='seconds')}`",
        f"- Goal: `{results.get('goal', 'N/A')}`",
        f"- Mode: `{results.get('mode', 'unknown')}`",
        f"- Success: `{success}`",
        f"- Best score: `{best_score:.2f}`",
        f"- Iterations: `{len(iterations)}`",
        f"- Duration (s): `{duration:.2f}`",
        "",
        "## Metrics",
        "",
        f"- Average score: `{avg_score:.2f}`",
        f"- Max iteration score: `{max_score:.2f}`",
        f"- Total queries: `{results.get('total_queries', 'N/A')}`",
        f"- Prompt chars: `{results.get('prompt_chars', 'N/A')}`",
        f"- Response chars: `{results.get('response_chars', 'N/A')}`",
        "",
        "## Distribution",
        "",
        f"- Min score: `{min_score:.2f}`",
        f"- Median score: `{med_score:.2f}`",
        f"- Average iteration time (s): `{avg_iter_sec:.2f}`",
        "",
        "## Top Sources",
        "",
    ]
    lines.extend([f"- `{name}`: `{count}`" for name, count in top_sources] or ["- `N/A`"])

    summary_path.write_text("\n".join(lines), encoding="utf-8")
    return summary_path


def _write_charts(output_dir: Path, iterations: List[Dict[str, Any]]) -> Dict[str, Path]:
    chart_paths: Dict[str, Path] = {}
    if not iterations:
        return chart_paths

    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return chart_paths

    plt.style.use("seaborn-v0_8-whitegrid")
    charts_dir = output_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)

    x = list(range(1, len(iterations) + 1))
    scores = [_to_float(it.get("score", 0.0)) for it in iterations]
    iter_secs = [_to_float(it.get("iteration_seconds", 0.0)) for it in iterations]
    sources = [str(it.get("source", "unknown")) for it in iterations]

    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(x, scores, marker="o", linewidth=1.8, color="#1f77b4")
    if len(scores) >= 2:
        running = [sum(scores[:i]) / i for i in range(1, len(scores) + 1)]
        ax.plot(x, running, linewidth=1.6, linestyle="--", color="#2ca02c", label="running mean")
        ax.legend(frameon=False)
    ax.set_title("Score Progression")
    ax.set_xlabel("Iteration")
    ax.set_ylabel("Score")
    ax.grid(alpha=0.3, linestyle=":")
    score_path = charts_dir / "score_progression.png"
    fig.savefig(score_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    chart_paths["score_progression"] = score_path

    if any(v > 0 for v in iter_secs):
        fig, ax = plt.subplots(figsize=(10, 5))
        ax.plot(x, iter_secs, marker="o", linewidth=1.8, color="#ff7f0e")
        ax.set_title("Iteration Timing")
        ax.set_xlabel("Iteration")
        ax.set_ylabel("Seconds")
        ax.grid(alpha=0.3, linestyle=":")
        time_path = charts_dir / "iteration_timing.png"
        fig.savefig(time_path, dpi=300, bbox_inches="tight")
        plt.close(fig)
        chart_paths["iteration_timing"] = time_path

    source_counts = Counter(sources)
    if source_counts:
        labels = list(source_counts.keys())
        values = list(source_counts.values())
        fig, ax = plt.subplots(figsize=(10, 5))
        ax.bar(labels, values, color="#2ca02c")
        ax.set_title("Source Distribution")
        ax.set_xlabel("Source")
        ax.set_ylabel("Count")
        ax.tick_params(axis="x", rotation=25)
        ax.grid(alpha=0.25, axis="y", linestyle=":")
        src_path = charts_dir / "source_distribution.png"
        fig.savefig(src_path, dpi=300, bbox_inches="tight")
        plt.close(fig)
        chart_paths["source_distribution"] = src_path

    # one-page dashboard
    fig, axes = plt.subplots(2, 2, figsize=(13, 8))
    axes = axes.flatten()
    axes[0].plot(x, scores, marker="o", linewidth=1.6, color="#1f77b4")
    axes[0].set_title("Score")
    axes[1].hist(scores, bins=min(12, max(4, len(scores) // 3)), color="#2ca02c", alpha=0.8)
    axes[1].set_title("Score Histogram")
    if any(v > 0 for v in iter_secs):
        axes[2].plot(x, iter_secs, marker="o", linewidth=1.6, color="#ff7f0e")
    axes[2].set_title("Iteration Seconds")
    src = Counter(sources)
    axes[3].bar(list(src.keys()), list(src.values()), color="#9467bd")
    axes[3].tick_params(axis="x", rotation=25)
    axes[3].set_title("Sources")
    for ax in axes:
        ax.grid(alpha=0.25, linestyle=":")
    fig.suptitle("Run Dashboard", fontsize=14, fontweight="bold")
    dash_path = charts_dir / "run_dashboard.png"
    fig.savefig(dash_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    chart_paths["run_dashboard"] = dash_path

    return chart_paths


def _write_excel(output_dir: Path, results: Dict[str, Any], iterations: List[Dict[str, Any]]) -> Path | None:
    try:
        import openpyxl
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:
        return None

    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    summary_rows = [
        ("Generated", datetime.now().isoformat(timespec="seconds")),
        ("Goal", str(results.get("goal", "N/A"))),
        ("Mode", str(results.get("mode", "unknown"))),
        ("Success", str(bool(results.get("success", False)))),
        ("Best score", _to_float(results.get("best_score", 0.0))),
        ("Iterations", len(iterations)),
        ("Duration (s)", _to_float(results.get("duration_seconds", 0.0))),
        ("Total queries", results.get("total_queries", "N/A")),
        ("Prompt chars", results.get("prompt_chars", "N/A")),
        ("Response chars", results.get("response_chars", "N/A")),
    ]

    ws_summary["A1"] = "Metric"
    ws_summary["B1"] = "Value"
    for cell in (ws_summary["A1"], ws_summary["B1"]):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, (k, v) in enumerate(summary_rows, start=2):
        ws_summary.cell(row=i, column=1, value=k)
        ws_summary.cell(row=i, column=2, value=v)
        ws_summary.cell(row=i, column=1).font = Font(bold=True)

    ws_summary.column_dimensions["A"].width = 24
    ws_summary.column_dimensions["B"].width = 70

    ws_iter = wb.create_sheet("Iterations")
    columns = [
        "iteration",
        "score",
        "success",
        "source",
        "prompt_len",
        "response_len",
        "iteration_seconds",
    ]
    for idx, col in enumerate(columns, start=1):
        cell = ws_iter.cell(row=1, column=idx, value=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, it in enumerate(iterations, start=2):
        row = {
            "iteration": it.get("iteration", i - 1),
            "score": _to_float(it.get("score", 0.0)),
            "success": bool(it.get("success", False)),
            "source": str(it.get("source", "unknown")),
            "prompt_len": len(str(it.get("prompt", ""))),
            "response_len": len(str(it.get("response", ""))),
            "iteration_seconds": _to_float(it.get("iteration_seconds", 0.0)),
        }
        for j, col in enumerate(columns, start=1):
            cell = ws_iter.cell(row=i, column=j, value=row[col])
            if col == "success":
                cell.alignment = Alignment(horizontal="center")

    widths = [12, 10, 10, 18, 12, 12, 16]
    for col, width in zip("ABCDEFG", widths):
        ws_iter.column_dimensions[col].width = width

    last_row = len(iterations) + 1
    if last_row >= 2:
        ws_iter.conditional_formatting.add(
            f"B2:B{last_row}",
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="F8696B",
                mid_type="num",
                mid_value=5,
                mid_color="FFEB84",
                end_type="num",
                end_value=10,
                end_color="63BE7B",
            ),
        )

    ws_kpi = wb.create_sheet("KPI")
    ws_kpi["A1"] = "KPI"
    ws_kpi["B1"] = "Value"
    ws_kpi["A1"].fill = header_fill
    ws_kpi["B1"].fill = header_fill
    ws_kpi["A1"].font = header_font
    ws_kpi["B1"].font = header_font
    scores = [_to_float(it.get("score", 0.0)) for it in iterations]
    kpis = [
        ("Best score", _to_float(results.get("best_score", 0.0))),
        ("Average score", round((sum(scores) / len(scores)) if scores else 0.0, 4)),
        ("Median score", round(median(scores), 4) if scores else 0.0),
        ("Score std proxy (range)", round((max(scores) - min(scores)) if scores else 0.0, 4)),
        ("Success", bool(results.get("success", False))),
    ]
    for i, (k, v) in enumerate(kpis, start=2):
        ws_kpi.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws_kpi.cell(row=i, column=2, value=v)
    ws_kpi.column_dimensions["A"].width = 30
    ws_kpi.column_dimensions["B"].width = 22

    xl_path = output_dir / "improved_results.xlsx"
    wb.save(xl_path)
    return xl_path


def generate_improved_outputs(results: Dict[str, Any], output_root: Path) -> Dict[str, Any]:
    """Generate improved text/image/excel outputs for a single run."""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = Path(output_root) / f"improved_outputs_{ts}"
    output_dir.mkdir(parents=True, exist_ok=True)

    iterations = _get_iterations(results)

    summary_path = _write_text_summary(output_dir, results, iterations)
    chart_paths = _write_charts(output_dir, iterations)
    excel_path = _write_excel(output_dir, results, iterations)

    manifest = {
        "output_dir": str(output_dir),
        "summary": str(summary_path),
        "charts": {k: str(v) for k, v in chart_paths.items()},
        "excel": str(excel_path) if excel_path else None,
        "iterations_count": len(iterations),
    }

    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    manifest["manifest"] = str(manifest_path)
    return manifest
